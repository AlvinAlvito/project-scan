"""Ambil kunci jawaban dari semua PDF "Pembahasan Soal" di satu folder ke Excel.

Pakai:
    python tools/import_kunci.py FOLDER_PDF [FILE_EXCEL]

Excel berisi satu baris per bidang (nama usulan, jumlah soal, kunci, catatan)
dan kolom Kode Naskah yang diisi panitia (kode naskah tidak tercantum di PDF).
Sheet kedua menampilkan kunci per nomor soal untuk dicocokkan dengan PDF.
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from ljk import kunci as KU  # noqa: E402

FILL_HEAD = PatternFill("solid", fgColor="D9D9D9")
FILL_ISI = PatternFill("solid", fgColor="FFF2CC")
FILL_MASALAH = PatternFill("solid", fgColor="FFD8A8")


def main(src, out=None):
    src = Path(src)
    pdfs = sorted(p for p in src.iterdir() if p.suffix.lower() == ".pdf")
    rows = [(KU.suggest_name(p), KU.parse(p)) for p in pdfs]
    rows.sort(key=lambda r: r[0])
    out = Path(out) if out else (Path(__file__).resolve().parents[1] / "output" / "kunci"
                                  / f"Kunci Jawaban - {src.name}.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Kunci"
    head = ["No", "Nama Bidang", "Kode Naskah", "Jumlah Soal", "Kunci", "Catatan", "File PDF"]
    widths = [5, 24, 13, 12, 80, 45, 60]
    for c, (h, w) in enumerate(zip(head, widths), 1):
        cell = ws.cell(1, c, h)
        cell.font, cell.fill = Font(bold=True), FILL_HEAD
        ws.column_dimensions[get_column_letter(c)].width = w
    for i, (name, k) in enumerate(rows, 1):
        vals = [i, name, "", k.jumlah_soal, k.kunci, "; ".join(k.masalah), str(k.path)]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(i + 1, c, v)
            if c == 5:
                cell.font = Font(name="Consolas")
        ws.cell(i + 1, 3).fill = FILL_ISI
        if k.masalah:
            ws.cell(i + 1, 6).fill = FILL_MASALAH
    ws.freeze_panes = "A2"

    nmax = max(k.jumlah_soal for _, k in rows)
    ws2 = wb.create_sheet("Per Nomor")
    ws2.cell(1, 1, "Nama Bidang").font = Font(bold=True)
    ws2.column_dimensions["A"].width = 24
    for q in range(1, nmax + 1):
        cell = ws2.cell(1, q + 1, q)
        cell.font, cell.fill = Font(bold=True), FILL_HEAD
        cell.alignment = Alignment(horizontal="center")
        ws2.column_dimensions[get_column_letter(q + 1)].width = 3.6
    for i, (name, k) in enumerate(rows, 2):
        ws2.cell(i, 1, name)
        for q, ch in enumerate(k.kunci, 2):
            ws2.cell(i, q, ch).alignment = Alignment(horizontal="center")
    ws2.freeze_panes = "B2"

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    for name, k in rows:
        print(f"{name:22s} {k.jumlah_soal:3d} soal  {'; '.join(k.masalah) or 'lengkap'}")
    print("Tersimpan:", out)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
