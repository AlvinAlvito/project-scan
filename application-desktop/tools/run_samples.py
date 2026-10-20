"""Baca semua foto LJK di satu folder dan simpan hasilnya untuk dicek.

Pakai:
    python tools/run_samples.py FOLDER_FOTO [FOLDER_HASIL]

Hasil:
    FOLDER_HASIL/Hasil Baca LJK.xlsx   ringkasan + jawaban per lembar
    FOLDER_HASIL/cek_visual/*.jpg      lembar yang sudah diluruskan + tanda hasil baca
    FOLDER_HASIL/potongan/<foto>/*.png potongan area tulisan tangan
Urutan lembar = urutan nama file (= urutan waktu foto diambil).
"""
import sys
import time
from pathlib import Path

import cv2
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from ljk import pipeline as PL  # noqa: E402
from ljk import template as T  # noqa: E402

EXTS = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff"}
FILL_RAGU = PatternFill("solid", fgColor="FFD8A8")
FILL_GANDA = PatternFill("solid", fgColor="FFB3B3")
FILL_CEK = PatternFill("solid", fgColor="FFF2CC")
FILL_GAGAL = PatternFill("solid", fgColor="FFB3B3")
FILL_HEAD = PatternFill("solid", fgColor="D9D9D9")


def write_png(path, img):
    path.parent.mkdir(parents=True, exist_ok=True)
    ok, buf = cv2.imencode(path.suffix, img)
    if ok:
        buf.tofile(str(path))


def header(ws, names, widths):
    for c, (n, w) in enumerate(zip(names, widths), 1):
        cell = ws.cell(1, c, n)
        cell.font = Font(bold=True)
        cell.fill = FILL_HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"


def main(src, out):
    src, out = Path(src), Path(out)
    files = sorted(p for p in src.iterdir() if p.suffix.lower() in EXTS)
    results = []
    for p in files:
        t = time.time()
        r = PL.process(p)
        results.append(r)
        if r.error:
            print(f"{p.name}: GAGAL - {r.error}")
            continue
        rd = r.read
        print(f"{p.name}: {r.status:9s} {rd.nama!r} no={rd.nomor} kode={rd.kode} tgl={rd.tanggal} "
              f"cek={len(r.issues)} ({time.time() - t:.1f}s)")
        write_png(out / "cek_visual" / f"{p.stem}.jpg", PL.overlay(r))
        for name, img in PL.handwriting_crops(r).items():
            write_png(out / "potongan" / p.stem / f"{name}.png", img)

    wb = Workbook()
    ws = wb.active
    ws.title = "Ringkasan"
    header(ws, ["No", "File", "Status", "Nama", "Nomor Peserta", "Kode Naskah", "Tanggal",
                "Soal terisi", "Hal yang perlu dicek"], [5, 26, 10, 24, 14, 12, 11, 10, 90])
    ws2 = wb.create_sheet("Jawaban")
    header(ws2, ["No", "File", "Nama", "Nomor"] + [str(q) for q in range(1, T.N_QUESTIONS + 1)],
           [5, 26, 24, 9] + [3.2] * T.N_QUESTIONS)
    for i, r in enumerate(results, 1):
        row = i + 1
        if r.error:
            vals = [i, r.path.name, r.status, "", "", "", "", "", r.error]
        else:
            rd = r.read
            terisi = sum(1 for a in rd.jawaban if a)
            vals = [i, r.path.name, r.status, rd.nama, rd.nomor, rd.kode, rd.tanggal, terisi,
                    "\n".join(r.issues)]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row, c, v)
            cell.alignment = Alignment(vertical="top", wrap_text=(c == 9))
        if r.status != "OK":
            ws.cell(row, 3).fill = FILL_GAGAL if r.error else FILL_CEK
        if r.error:
            continue
        for c, v in enumerate([i, r.path.name, r.read.nama, r.read.nomor], 1):
            ws2.cell(row, c, v)
        for q, g in enumerate(r.read.answer_groups):
            cell = ws2.cell(row, 5 + q, g.value)
            cell.alignment = Alignment(horizontal="center")
            if g.kind == "ragu":
                cell.fill = FILL_RAGU
            elif g.kind == "ganda":
                cell.fill = FILL_GANDA
    out.mkdir(parents=True, exist_ok=True)
    xlsx = out / "Hasil Baca LJK.xlsx"
    wb.save(xlsx)
    n_ok = sum(r.status == "OK" for r in results)
    n_cek = sum(r.status == "Perlu cek" for r in results)
    n_fail = sum(r.status == "Gagal" for r in results)
    print(f"\nSelesai: {len(results)} lembar | OK {n_ok} | Perlu cek {n_cek} | Gagal {n_fail}")
    print("Excel:", xlsx)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    default_out = Path(__file__).resolve().parents[1] / "output" / "sampel"
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else default_out)
