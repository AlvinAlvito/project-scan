"""Rekap satu event: baca semua foto LJK, nilai dengan kunci, ekspor Daftar Nilai.

Pakai:
    python tools/rekap.py FOLDER_FOTO FILE_KUNCI.xlsx --event "NAMA EVENT" --tanggal "6 September 2026" [--out FOLDER]

Hasil:
    <event> Daftar Nilai.xlsx   satu sheet per bidang (format yang disepakati)
    <event> Perlu Cek.xlsx      lembar/soal yang harus dicek manusia sebelum nilai final
Lembar dengan kode naskah yang tidak terbaca/tidak terdaftar TIDAK masuk Daftar
Nilai (tidak bisa dinilai), hanya muncul di Perlu Cek.
"""
import argparse
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from ljk import export as EX  # noqa: E402
from ljk import nilai as NI  # noqa: E402
from ljk import pipeline as PL  # noqa: E402

EXTS = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff"}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("foto")
    ap.add_argument("kunci")
    ap.add_argument("--event", required=True)
    ap.add_argument("--tanggal", required=True)
    ap.add_argument("--out", default=str(Path(__file__).resolve().parents[1] / "output"))
    a = ap.parse_args()

    soal = NI.load_kunci(a.kunci)
    files = sorted(p for p in Path(a.foto).iterdir() if p.suffix.lower() in EXTS)
    hasil = []           # (urutan, file, SheetResult, Soal|None, Skor|None, [catatan])
    for i, p in enumerate(files, 1):
        r = PL.process(p)
        notes = list(r.issues) if not r.error else [r.error]
        s = sk = None
        if not r.error:
            s = soal.get(r.read.kode)
            if s is None:
                notes.insert(0, f"kode naskah '{r.read.kode}' tidak terbaca/tidak terdaftar, lembar belum dinilai")
            else:
                sk = NI.nilai(r.read.jawaban, s)
                notes += sk.catatan
        hasil.append((i, p.name, r, s, sk, notes))
        print(f"{p.name}: {'GAGAL' if r.error else (s.bidang if s else 'kode?')} "
              f"{'' if sk is None else f'B={sk.benar} S={sk.salah} K={sk.kosong} nilai={sk.nilai}'}")

    # Nomor peserta ganda dalam satu event.
    nomor = Counter(h[2].read.nomor for h in hasil if not h[2].error)
    for h in hasil:
        if not h[2].error and nomor[h[2].read.nomor] > 1:
            h[5].append(f"nomor peserta {h[2].read.nomor} dipakai lebih dari satu lembar")

    groups = {}
    for i, fname, r, s, sk, _ in hasil:
        if sk is not None:
            groups.setdefault(s.kode, (s.bidang, []))[1].append((r.read.nomor, r.read.nama, sk))
    ordered = [groups[k] for k in sorted(groups, key=lambda x: (len(x), x))]

    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    dn = out / f"{a.event} Daftar Nilai.xlsx"
    EX.daftar_nilai(dn, a.event, a.tanggal, ordered)

    wb = Workbook()
    ws = wb.active
    ws.title = "Perlu Cek"
    head = ["Urutan", "File", "Nama", "Nomor", "Kode", "Bidang", "Nilai (sementara)", "Hal yang perlu dicek"]
    widths = [8, 26, 24, 9, 7, 20, 10, 110]
    for c, (h, w) in enumerate(zip(head, widths), 1):
        cell = ws.cell(1, c, h)
        cell.font, cell.fill = Font(bold=True), PatternFill("solid", fgColor="D9D9D9")
        ws.column_dimensions[get_column_letter(c)].width = w
    row = 2
    for i, fname, r, s, sk, notes in hasil:
        if not notes:
            continue
        rd = r.read
        vals = [i, fname, rd.nama if rd else "", rd.nomor if rd else "", rd.kode if rd else "",
                s.bidang if s else "", sk.nilai if sk else "", "\n".join(notes)]
        for c, v in enumerate(vals, 1):
            ws.cell(row, c, v).alignment = Alignment(vertical="top", wrap_text=(c == 8))
        row += 1
    ws.freeze_panes = "A2"
    pc = out / f"{a.event} Perlu Cek.xlsx"
    wb.save(pc)

    n_cek = row - 2
    print(f"\n{len(hasil)} lembar | dinilai {sum(1 for h in hasil if h[4])} | perlu cek {n_cek}")
    print("Daftar Nilai:", dn)
    print("Perlu Cek   :", pc)


if __name__ == "__main__":
    main()
