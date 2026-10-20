"""Ekspor Daftar Nilai ke Excel, mengikuti format contoh FSN KEPRI.

Perbedaan yang sudah disepakati dengan contoh: Nomor Peserta ditulis polos
(mis. 10277), tanpa kolom Keterangan, tanpa kalimat batas remedial, dan tanpa
blok tanda tangan. Satu sheet per nama bidang; urutan peserta = urutan scan.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="C8C8C8")
COLS = ["No.", "Nomor Peserta", "Nama", "B", "S", "K", "Nilai"]
WIDTHS = {"A": 2.5, "B": 6.6, "C": 14.5, "D": 27.6, "E": 6.6, "F": 6.6, "G": 6.6, "H": 7.6}


def _sheet_title(name, used):
    t = "".join(ch for ch in name if ch not in '[]:*?/\\')[:31] or "BIDANG"
    base, n = t, 2
    while t in used:
        t = f"{base[:28]} {n}"
        n += 1
    used.add(t)
    return t


def daftar_nilai(path, event, tanggal, groups):
    """groups: [(nama_bidang, [(nomor, nama, Skor), ...]), ...] sesuai urutan sheet."""
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for bidang, rows in groups:
        ws = wb.create_sheet(_sheet_title(bidang, used))
        for col, w in WIDTHS.items():
            ws.column_dimensions[col].width = w
        titles = [("DAFTAR NILAI", 14), ("POSI Indonesia", 14), (f"{event} - {bidang} ({tanggal})", 10)]
        for r, (text, size) in enumerate(titles, 2):
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
            c = ws.cell(r, 2, text)
            c.font = Font(name="Arial", size=size, bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = ws.row_dimensions[3].height = 17.65
        ws.row_dimensions[5].height = 21.5
        for c, h in enumerate(COLS, 2):
            cell = ws.cell(5, c, h)
            cell.font = Font(name="Arial", size=11, bold=True)
            cell.fill = HEAD_FILL
            cell.border = BOX
            cell.alignment = Alignment(horizontal="center" if h in ("Nomor Peserta", "Nama") else "right",
                                       vertical="center")
        r = 6
        for i, (nomor, nama, sk) in enumerate(rows, 1):
            vals = [i, nomor, nama, sk.benar, sk.salah, sk.kosong, sk.nilai]
            for c, v in enumerate(vals, 2):
                cell = ws.cell(r, c, v)
                cell.font = Font(name="Arial", size=11)
                cell.border = BOX
            ws.cell(r, 8).number_format = "0.00"
            r += 1
        if rows:
            ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
            ws.cell(r, 3, "Rata-Rata").font = Font(name="Arial", size=10)
            ws.cell(r, 3).alignment = Alignment(vertical="center")
            n = len(rows)
            avgs = [sum(x[2].benar for x in rows) / n, sum(x[2].salah for x in rows) / n,
                    sum(x[2].kosong for x in rows) / n, sum(x[2].nilai for x in rows) / n]
            for c, v in enumerate(avgs, 5):
                cell = ws.cell(r, c, v)
                cell.font = Font(name="Arial", size=11)
                cell.number_format = "0.00"
            for c in range(2, 9):
                ws.cell(r, c).border = BOX
    wb.save(path)
