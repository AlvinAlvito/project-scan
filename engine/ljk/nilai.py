"""Kunci jawaban per kode naskah dan penilaian.

Aturan nilai (keputusan panitia): benar +4, salah -1, kosong 0.
Lebih dari satu pilihan dihitamkan (ganda) dihitung SALAH.
Bulatan "ragu" yang belum dicek manusia sementara dihitung kosong dan
selalu dilaporkan di daftar "perlu cek".
"""
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

BENAR, SALAH, KOSONG = 4, -1, 0


@dataclass
class Soal:
    kode: str
    bidang: str
    kunci: str
    sumber: str = ""        # path PDF pembahasan asal kunci (bila diimpor dari PDF)

    @property
    def jumlah(self):
        return len(self.kunci)


def load_kunci(xlsx):
    """Baca file kunci (hasil tools/import_kunci.py yang Kode Naskah-nya sudah diisi).

    Mengembalikan {kode_naskah: Soal}."""
    ws = load_workbook(xlsx, read_only=True)["Kunci"]
    rows = ws.iter_rows(values_only=True)
    head = [str(h).strip() if h else "" for h in next(rows)]
    ix = {h: i for i, h in enumerate(head)}
    out = {}
    for r in rows:
        kode = r[ix["Kode Naskah"]]
        if kode in (None, ""):
            continue
        kode = str(kode).strip()
        if kode.endswith(".0"):
            kode = kode[:-2]
        kunci = str(r[ix["Kunci"]] or "").strip().upper()
        if kode in out:
            raise ValueError(f"Kode naskah {kode} dipakai dua kali di file kunci")
        sumber = str(r[ix["File PDF"]] or "").strip() if "File PDF" in ix else ""
        out[kode] = Soal(kode, str(r[ix["Nama Bidang"]] or "").strip(), kunci, sumber)
    return out


@dataclass
class Skor:
    benar: int
    salah: int
    kosong: int
    catatan: list = field(default_factory=list)
    ragu: list = field(default_factory=list)     # nomor soal berisi "?"
    ganda: list = field(default_factory=list)    # nomor soal berisi "*"
    lebih: list = field(default_factory=list)    # nomor soal terisi setelah soal terakhir

    @property
    def nilai(self):
        return BENAR * self.benar + SALAH * self.salah + KOSONG * self.kosong


def nilai(jawaban, soal):
    """jawaban: daftar 180 isian ('A'..'E', '' kosong, '?' ragu, '*' ganda)."""
    b = s = k = 0
    ragu, ganda = [], []
    for q, kunci in enumerate(soal.kunci):
        a = jawaban[q]
        if a == "*":
            s += 1
            ganda.append(q + 1)
        elif a == "?":
            k += 1
            ragu.append(q + 1)
        elif a == "":
            k += 1
        elif a == kunci:
            b += 1
        else:
            s += 1
    cat = []
    if ganda:
        cat.append(f"jawaban ganda (dihitung salah): soal {', '.join(map(str, ganda))}")
    if ragu:
        cat.append(f"ragu, sementara dihitung kosong: soal {', '.join(map(str, ragu))}")
    # Tanda "kode naskah mungkin salah": peserta menjawab banyak soal TEPAT sesudah soal
    # terakhir ujian. Bercak di pojok bawah kertas (bayangan/terpotong) tidak dihitung.
    batas = min(soal.jumlah + 30, len(jawaban))
    lebih = [q + 1 for q in range(soal.jumlah, batas) if jawaban[q] not in ("", "?")]
    if len(lebih) < 3:
        lebih = []
    if lebih:
        cat.append(f"ada isian setelah soal terakhir ({soal.jumlah}): soal {lebih[0]}-{lebih[-1]}, "
                   "cek apakah kode naskah benar")
    return Skor(b, s, k, cat, ragu, ganda, lebih)
