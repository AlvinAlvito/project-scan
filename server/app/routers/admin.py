"""Panel superadmin: kompetisi, lokasi, bidang, peserta, akun & hak akses, dashboard, lembar."""
import csv
import io
from datetime import datetime
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import BaseModel, Field
from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from .. import models as M
from .. import reading
from ..db import get_db
from ..security import hash_password, superadmin
from ..serial import bidang_out, kompetisi_out, lembar_out, lokasi_out, peserta_out, user_out
from ..validate import validate

router = APIRouter(prefix="/api/admin", tags=["superadmin"], dependencies=[Depends(superadmin)])


def _get(db, model, id_, what):
    obj = db.get(model, id_)
    if obj is None:
        raise HTTPException(404, f"{what} tidak ditemukan.")
    return obj


def _commit(db, conflict_msg):
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, conflict_msg)


# ---------------------------------------------------------------- kompetisi

class KompetisiIn(BaseModel):
    nama: str = Field(min_length=1, max_length=160)
    tanggal: str = ""
    status: str = "aktif"


@router.get("/kompetisi")
def list_kompetisi(db: Session = Depends(get_db)):
    rows = db.scalars(select(M.Kompetisi).order_by(M.Kompetisi.id.desc())).all()
    counts = dict(db.execute(select(M.Lembar.kompetisi_id, func.count()).where(M.Lembar.status != M.GAGAL)
                             .group_by(M.Lembar.kompetisi_id)).all())
    out = []
    for k in rows:
        d = kompetisi_out(k, detail=True)
        d["jumlah_lembar"] = counts.get(k.id, 0)
        out.append(d)
    return out


@router.post("/kompetisi")
def create_kompetisi(body: KompetisiIn, db: Session = Depends(get_db)):
    k = M.Kompetisi(nama=body.nama.strip(), tanggal=body.tanggal.strip(), status=body.status)
    db.add(k)
    db.commit()
    return kompetisi_out(k, detail=True)


@router.get("/kompetisi/{kid}")
def get_kompetisi(kid: int, db: Session = Depends(get_db)):
    return kompetisi_out(_get(db, M.Kompetisi, kid, "Kompetisi"), detail=True)


@router.put("/kompetisi/{kid}")
def update_kompetisi(kid: int, body: KompetisiIn, db: Session = Depends(get_db)):
    k = _get(db, M.Kompetisi, kid, "Kompetisi")
    k.nama, k.tanggal, k.status = body.nama.strip(), body.tanggal.strip(), body.status
    db.commit()
    return kompetisi_out(k, detail=True)


@router.delete("/kompetisi/{kid}")
def delete_kompetisi(kid: int, db: Session = Depends(get_db)):
    k = _get(db, M.Kompetisi, kid, "Kompetisi")
    n = db.scalar(select(func.count()).select_from(M.Lembar).where(M.Lembar.kompetisi_id == kid))
    if n:
        raise HTTPException(409, f"Kompetisi ini sudah punya {n} lembar. Ubah statusnya menjadi selesai "
                                 "daripada menghapus.")
    db.delete(k)
    db.commit()
    return {"ok": True}


# ---------------------------------------------------------------- lokasi

class LokasiIn(BaseModel):
    kota: str = Field(min_length=1, max_length=120)
    target_lembar: int = 0


@router.post("/kompetisi/{kid}/lokasi")
def create_lokasi(kid: int, body: LokasiIn, db: Session = Depends(get_db)):
    _get(db, M.Kompetisi, kid, "Kompetisi")
    lo = M.Lokasi(kompetisi_id=kid, kota=body.kota.strip(), target_lembar=max(0, body.target_lembar))
    db.add(lo)
    db.commit()
    return lokasi_out(lo)


@router.put("/lokasi/{lid}")
def update_lokasi(lid: int, body: LokasiIn, db: Session = Depends(get_db)):
    lo = _get(db, M.Lokasi, lid, "Lokasi")
    lo.kota, lo.target_lembar = body.kota.strip(), max(0, body.target_lembar)
    db.commit()
    return lokasi_out(lo)


@router.delete("/lokasi/{lid}")
def delete_lokasi(lid: int, db: Session = Depends(get_db)):
    lo = _get(db, M.Lokasi, lid, "Lokasi")
    n = db.scalar(select(func.count()).select_from(M.Lembar).where(M.Lembar.lokasi_id == lid))
    if n:
        raise HTTPException(409, f"Lokasi ini sudah punya {n} lembar dan tidak bisa dihapus.")
    db.delete(lo)
    db.commit()
    return {"ok": True}


# ---------------------------------------------------------------- bidang

class BidangIn(BaseModel):
    kode_naskah: str = Field(min_length=1, max_length=10)
    nama: str = Field(min_length=1, max_length=120)
    jumlah_soal: int = Field(ge=1, le=180)


@router.post("/kompetisi/{kid}/bidang")
def create_bidang(kid: int, body: BidangIn, db: Session = Depends(get_db)):
    _get(db, M.Kompetisi, kid, "Kompetisi")
    b = M.Bidang(kompetisi_id=kid, kode_naskah=body.kode_naskah.strip(), nama=body.nama.strip(),
                 jumlah_soal=body.jumlah_soal)
    db.add(b)
    _commit(db, f"Kode naskah {body.kode_naskah} sudah dipakai di kompetisi ini.")
    return bidang_out(b)


@router.put("/bidang/{bid}")
def update_bidang(bid: int, body: BidangIn, db: Session = Depends(get_db)):
    b = _get(db, M.Bidang, bid, "Bidang")
    b.kode_naskah, b.nama, b.jumlah_soal = body.kode_naskah.strip(), body.nama.strip(), body.jumlah_soal
    _commit(db, f"Kode naskah {body.kode_naskah} sudah dipakai di kompetisi ini.")
    return bidang_out(b)


@router.delete("/bidang/{bid}")
def delete_bidang(bid: int, db: Session = Depends(get_db)):
    b = _get(db, M.Bidang, bid, "Bidang")
    n = db.scalar(select(func.count()).select_from(M.Lembar).where(M.Lembar.bidang_id == bid))
    if n:
        raise HTTPException(409, f"Bidang ini sudah punya {n} lembar dan tidak bisa dihapus.")
    db.delete(b)
    db.commit()
    return {"ok": True}


# ---------------------------------------------------------------- master bidang

class MasterBidangIn(BaseModel):
    nama: str = Field(min_length=1, max_length=120)
    kode_naskah: str = Field("", max_length=10)
    jumlah_soal: int = Field(0, ge=0, le=180)


def _master_out(m: M.MasterBidang, dipakai=0):
    return {"id": m.id, "nama": m.nama, "kode_naskah": m.kode_naskah, "jumlah_soal": m.jumlah_soal,
            "dipakai": dipakai}


@router.get("/master-bidang")
def list_master_bidang(db: Session = Depends(get_db)):
    pakai = dict(db.execute(select(M.Bidang.nama, func.count()).group_by(M.Bidang.nama)).all())
    return [_master_out(m, pakai.get(m.nama, 0))
            for m in db.scalars(select(M.MasterBidang).order_by(M.MasterBidang.nama))]


@router.post("/master-bidang")
def create_master_bidang(body: MasterBidangIn, db: Session = Depends(get_db)):
    m = M.MasterBidang(nama=" ".join(body.nama.split()), kode_naskah=body.kode_naskah.strip(),
                       jumlah_soal=body.jumlah_soal)
    db.add(m)
    _commit(db, f"Bidang '{body.nama}' sudah ada di master.")
    return _master_out(m)


@router.put("/master-bidang/{mid}")
def update_master_bidang(mid: int, body: MasterBidangIn, db: Session = Depends(get_db)):
    m = _get(db, M.MasterBidang, mid, "Bidang")
    m.nama, m.kode_naskah, m.jumlah_soal = " ".join(body.nama.split()), body.kode_naskah.strip(), body.jumlah_soal
    _commit(db, f"Bidang '{body.nama}' sudah ada di master.")
    return _master_out(m)


@router.delete("/master-bidang/{mid}")
def delete_master_bidang(mid: int, db: Session = Depends(get_db)):
    """Bidang yang sudah dipakai kompetisi tetap ada di kompetisi itu; hanya hilang dari pilihan."""
    db.delete(_get(db, M.MasterBidang, mid, "Bidang"))
    db.commit()
    return {"ok": True}


# ---------------------------------------------------------------- peserta

class PesertaIn(BaseModel):
    nomor: str = Field(min_length=1, max_length=30)
    nama: str = Field(min_length=1, max_length=160)
    sekolah: str = ""
    lokasi_id: int | None = None


@router.get("/kompetisi/{kid}/peserta")
def list_peserta(kid: int, q: str = "", lokasi_id: int | None = None, limit: int = Query(100, le=1000),
                 offset: int = 0, db: Session = Depends(get_db)):
    cond = [M.Peserta.kompetisi_id == kid]
    if q:
        like = f"%{q.strip()}%"
        cond.append(or_(M.Peserta.nomor.like(like), M.Peserta.nama.like(like), M.Peserta.sekolah.like(like)))
    if lokasi_id:
        cond.append(M.Peserta.lokasi_id == lokasi_id)
    total = db.scalar(select(func.count()).select_from(M.Peserta).where(*cond))
    rows = db.scalars(select(M.Peserta).where(*cond).order_by(M.Peserta.nomor).limit(limit).offset(offset)).all()
    return {"total": total, "items": [peserta_out(p) for p in rows]}


@router.post("/kompetisi/{kid}/peserta")
def create_peserta(kid: int, body: PesertaIn, db: Session = Depends(get_db)):
    _get(db, M.Kompetisi, kid, "Kompetisi")
    p = M.Peserta(kompetisi_id=kid, nomor=body.nomor.strip(), nama=" ".join(body.nama.upper().split()),
                  sekolah=body.sekolah.strip(), lokasi_id=body.lokasi_id)
    db.add(p)
    _commit(db, f"Nomor peserta {body.nomor} sudah terdaftar.")
    return peserta_out(p)


@router.put("/peserta/{pid}")
def update_peserta(pid: int, body: PesertaIn, db: Session = Depends(get_db)):
    p = _get(db, M.Peserta, pid, "Peserta")
    p.nomor, p.nama, p.sekolah, p.lokasi_id = (body.nomor.strip(), " ".join(body.nama.upper().split()),
                                               body.sekolah.strip(), body.lokasi_id)
    _commit(db, f"Nomor peserta {body.nomor} sudah terdaftar.")
    return peserta_out(p)


@router.delete("/peserta/{pid}")
def delete_peserta(pid: int, db: Session = Depends(get_db)):
    db.delete(_get(db, M.Peserta, pid, "Peserta"))
    db.commit()
    return {"ok": True}


def _rows_from_upload(name: str, data: bytes):
    if name.lower().endswith(".xlsx"):
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        return [[("" if c is None else str(c)).strip() for c in r] for r in ws.iter_rows(values_only=True)]
    text = data.decode("utf-8-sig", errors="replace")
    dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t") if text.strip() else csv.excel
    return [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), dialect)]


@router.post("/kompetisi/{kid}/peserta/import")
def import_peserta(kid: int, file: UploadFile = File(...), db: Session = Depends(get_db)):
    """Kolom yang dikenali (baris pertama = judul): nomor, nama, sekolah, lokasi/kota."""
    k = _get(db, M.Kompetisi, kid, "Kompetisi")
    rows = _rows_from_upload(file.filename or "", file.file.read())
    if not rows:
        raise HTTPException(400, "File kosong.")
    head = [h.lower() for h in rows[0]]

    def col(*names):
        for i, h in enumerate(head):
            if any(n in h for n in names):
                return i
        return None
    c_no, c_nama, c_sek, c_lok = col("nomor", "no peserta", "no. peserta"), col("nama"), col("sekolah"), \
        col("lokasi", "kota")
    if c_no is None or c_nama is None:
        raise HTTPException(400, "Baris pertama harus berisi judul kolom, minimal 'nomor' dan 'nama'.")
    kota = {lo.kota.strip().lower(): lo.id for lo in k.lokasi}
    existing = {p.nomor: p for p in db.scalars(select(M.Peserta).where(M.Peserta.kompetisi_id == kid))}
    baru = ubah = lewat = 0
    tanpa_lokasi = set()
    for r in rows[1:]:
        if not any(r):
            continue                    # baris kosong (mis. sisa format template) tidak dihitung
        get = lambda i: (r[i] if i is not None and i < len(r) else "").strip()  # noqa: E731
        nomor, nama = get(c_no), " ".join(get(c_nama).upper().split())
        if nomor.endswith(".0"):
            nomor = nomor[:-2]        # angka dari Excel
        if not nomor or not nama:
            lewat += 1
            continue
        lok_id = None
        if c_lok is not None and get(c_lok):
            lok_id = kota.get(get(c_lok).lower())
            if lok_id is None:
                tanpa_lokasi.add(get(c_lok))
        p = existing.get(nomor)
        if p is None:
            p = M.Peserta(kompetisi_id=kid, nomor=nomor, nama=nama, sekolah=get(c_sek), lokasi_id=lok_id)
            db.add(p)
            existing[nomor] = p
            baru += 1
        else:
            p.nama, p.sekolah = nama, get(c_sek) or p.sekolah
            p.lokasi_id = lok_id or p.lokasi_id
            ubah += 1
    db.commit()
    return {"baru": baru, "diperbarui": ubah, "dilewati": lewat, "lokasi_tidak_dikenal": sorted(tanpa_lokasi)}


@router.get("/kompetisi/{kid}/peserta/template")
def peserta_template(kid: int, db: Session = Depends(get_db)):
    """Template Excel yang pasti terbaca oleh impor peserta, berisi lokasi kompetisi ini."""
    k = _get(db, M.Kompetisi, kid, "Kompetisi")
    kota = [lo.kota for lo in k.lokasi] or ["(isi nama lokasi)"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Peserta"
    head_fill = PatternFill("solid", fgColor="1E3A8A")
    ws.append(["nomor", "nama", "sekolah", "lokasi"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center")
    contoh = [("10001", "BUDI SANTOSO", "SMA NEGERI 1 CONTOH"), ("10002", "SITI RAHMAWATI", "SMA NEGERI 2 CONTOH"),
              ("10003", "ANDI PRATAMA", "MAN 1 CONTOH")]      # data fiktif
    for i, (no, nama, sek) in enumerate(contoh):
        ws.append([no, nama, sek, kota[i % len(kota)]])
    for col, w in zip("ABCD", (14, 34, 34, 20)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, max_row=1000, min_col=1, max_col=1):
        for c in row:
            c.number_format = "@"          # teks, supaya nol di depan nomor tidak hilang
    ws.freeze_panes = "A2"
    if k.lokasi:
        dv = DataValidation(type="list", formula1="=Lokasi!$A$2:$A$" + str(len(kota) + 1), allow_blank=True,
                            showErrorMessage=True, errorTitle="Lokasi tidak dikenal",
                            error="Pilih lokasi dari daftar (sheet Lokasi).")
        ws.add_data_validation(dv)
        dv.add("D2:D1000")
    lw = wb.create_sheet("Lokasi")
    lw.append(["lokasi yang valid"])
    lw["A1"].font = Font(bold=True)
    for kt in kota:
        lw.append([kt])
    lw.column_dimensions["A"].width = 28
    pw = wb.create_sheet("Petunjuk")
    for line in (
        f"Template daftar peserta - {k.nama}",
        "",
        "1. Isi sheet 'Peserta' mulai baris 2. Baris contoh boleh dihapus atau ditimpa.",
        "2. Judul kolom di baris 1 jangan diubah: nomor, nama, sekolah, lokasi.",
        "3. nomor dan nama wajib diisi. Nomor ditulis persis seperti di LJK (kolom berformat teks).",
        "4. lokasi harus sama dengan salah satu nama di sheet 'Lokasi'; boleh dikosongkan.",
        "5. Nomor yang sudah terdaftar akan diperbarui datanya, bukan digandakan.",
        "6. Simpan sebagai .xlsx lalu impor lewat tombol 'Impor Excel/CSV'.",
    ):
        pw.append([line])
    pw["A1"].font = Font(bold=True, size=13)
    pw.column_dimensions["A"].width = 95
    buf = io.BytesIO()
    wb.save(buf)
    name = "Template Peserta - " + "".join(ch for ch in k.nama if ch.isalnum() or ch in " -_").strip() + ".xlsx"
    return Response(buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}"})


# ---------------------------------------------------------------- akun

class UserIn(BaseModel):
    nama: str = Field(min_length=1, max_length=120)
    username: str = Field(min_length=3, max_length=60)
    password: str | None = None
    peran: str = M.ADMIN_DAERAH
    aktif: bool = True
    lokasi_ids: list[int] = []


@router.get("/users")
def list_users(db: Session = Depends(get_db)):
    users = db.scalars(select(M.User).order_by(M.User.peran.desc(), M.User.nama)).all()
    return [user_out(u, with_akses=True) for u in users]


def _apply_user(db, u: M.User, body: UserIn):
    if body.peran not in (M.SUPERADMIN, M.ADMIN_DAERAH):
        raise HTTPException(400, "Peran tidak dikenal.")
    u.nama, u.username, u.peran, u.aktif = body.nama.strip(), body.username.strip().lower(), body.peran, body.aktif
    if body.password:
        if len(body.password) < 6:
            raise HTTPException(400, "Kata sandi minimal 6 karakter.")
        u.password_hash = hash_password(body.password)
    want = set(body.lokasi_ids) if body.peran == M.ADMIN_DAERAH else set()
    valid = set(db.scalars(select(M.Lokasi.id).where(M.Lokasi.id.in_(want)))) if want else set()
    u.akses = [a for a in u.akses if a.lokasi_id in valid]
    have = {a.lokasi_id for a in u.akses}
    u.akses += [M.Akses(lokasi_id=lid) for lid in valid - have]


@router.post("/users")
def create_user(body: UserIn, db: Session = Depends(get_db)):
    if not body.password:
        raise HTTPException(400, "Kata sandi wajib diisi untuk akun baru.")
    u = M.User(password_hash="")
    _apply_user(db, u, body)
    db.add(u)
    _commit(db, f"Username {body.username} sudah dipakai.")
    return user_out(u, with_akses=True)


@router.put("/users/{uid}")
def update_user(uid: int, body: UserIn, me: M.User = Depends(superadmin), db: Session = Depends(get_db)):
    u = _get(db, M.User, uid, "Akun")
    if u.id == me.id and (body.peran != M.SUPERADMIN or not body.aktif):
        raise HTTPException(400, "Anda tidak bisa menurunkan peran atau menonaktifkan akun sendiri.")
    _apply_user(db, u, body)
    _commit(db, f"Username {body.username} sudah dipakai.")
    return user_out(u, with_akses=True)


@router.delete("/users/{uid}")
def delete_user(uid: int, me: M.User = Depends(superadmin), db: Session = Depends(get_db)):
    if uid == me.id:
        raise HTTPException(400, "Anda tidak bisa menghapus akun sendiri.")
    db.delete(_get(db, M.User, uid, "Akun"))
    db.commit()
    return {"ok": True}


# ---------------------------------------------------------------- dashboard

@router.get("/dashboard/{kid}")
def dashboard(kid: int, db: Session = Depends(get_db)):
    k = _get(db, M.Kompetisi, kid, "Kompetisi")
    by = {}
    for lid, status, ditarik, n in db.execute(
            select(M.Lembar.lokasi_id, M.Lembar.status, M.Lembar.ditarik_at.is_not(None), func.count())
            .where(M.Lembar.kompetisi_id == kid).group_by(M.Lembar.lokasi_id, M.Lembar.status,
                                                          M.Lembar.ditarik_at.is_not(None))):
        d = by.setdefault(lid, {s: 0 for s in M.STATUS} | {"ditarik": 0})
        d[status] += n
        if ditarik:
            d["ditarik"] += n
    peserta = dict(db.execute(select(M.Peserta.lokasi_id, func.count()).where(M.Peserta.kompetisi_id == kid)
                              .group_by(M.Peserta.lokasi_id)).all())
    terakhir = dict(db.execute(select(M.Lembar.lokasi_id, func.max(M.Lembar.dibuat))
                               .where(M.Lembar.kompetisi_id == kid).group_by(M.Lembar.lokasi_id)).all())
    admins = {}
    for lid, nama in db.execute(select(M.Akses.lokasi_id, M.User.nama).join(M.User)
                                .where(M.Akses.lokasi_id.in_([lo.id for lo in k.lokasi]))):
        admins.setdefault(lid, []).append(nama)
    rows = []
    for lo in k.lokasi:
        d = by.get(lo.id, {s: 0 for s in M.STATUS} | {"ditarik": 0})
        terkirim = d[M.MENUNGGU] + d[M.DIKONFIRMASI] + d[M.DITANDAI]
        target = lo.target_lembar or peserta.get(lo.id, 0)
        rows.append({"lokasi": lokasi_out(lo), "target": target, "peserta": peserta.get(lo.id, 0),
                     "terkirim": terkirim, **d, "admin": admins.get(lo.id, []),
                     "terakhir": terakhir[lo.id].isoformat(sep=" ", timespec="minutes") if lo.id in terakhir else None})
    persetujuan = db.scalar(select(func.count()).select_from(M.Lembar).where(
        M.Lembar.kompetisi_id == kid, M.Lembar.peserta_baru_status == "diajukan"))
    return {"kompetisi": kompetisi_out(k), "lokasi": rows, "persetujuan": persetujuan or 0}


# ---------------------------------------------------------------- lembar

@router.get("/lembar")
def list_lembar(kompetisi_id: int, lokasi_id: int | None = None, status: str | None = None, q: str = "",
                peserta_baru: bool = False, limit: int = Query(50, le=200), offset: int = 0,
                db: Session = Depends(get_db)):
    cond = [M.Lembar.kompetisi_id == kompetisi_id]
    if lokasi_id:
        cond.append(M.Lembar.lokasi_id == lokasi_id)
    if status == "ditarik":
        cond.append(M.Lembar.ditarik_at.is_not(None))
    elif status:
        cond.append(M.Lembar.status == status)
    if peserta_baru:
        cond.append(M.Lembar.peserta_baru_status == "diajukan")
    if q:
        like = f"%{q.strip()}%"
        cond.append(or_(M.Lembar.hasil_baca["nomor"].as_string().like(like),
                        M.Lembar.hasil_baca["nama"].as_string().like(like)))
    total = db.scalar(select(func.count()).select_from(M.Lembar).where(*cond))
    rows = db.scalars(select(M.Lembar).where(*cond).order_by(M.Lembar.id.desc()).limit(limit).offset(offset)).all()
    return {"total": total, "items": [lembar_out(lb) for lb in rows]}


@router.get("/lembar/{lid}")
def get_lembar(lid: int, db: Session = Depends(get_db)):
    return lembar_out(_get(db, M.Lembar, lid, "Lembar"), detail=True)


@router.post("/lembar/{lid}/setujui-peserta")
def approve_peserta(lid: int, db: Session = Depends(get_db)):
    lb = _get(db, M.Lembar, lid, "Lembar")
    pb = lb.peserta_baru
    if not pb:
        raise HTTPException(400, "Lembar ini tidak mengajukan peserta baru.")
    p = db.scalar(select(M.Peserta).where(M.Peserta.kompetisi_id == lb.kompetisi_id, M.Peserta.nomor == pb["nomor"]))
    if p is None:
        p = M.Peserta(kompetisi_id=lb.kompetisi_id, lokasi_id=lb.lokasi_id, nomor=pb["nomor"],
                      nama=" ".join(pb.get("nama", "").upper().split()), sekolah=pb.get("sekolah", ""),
                      daftar_di_tempat=True)
        db.add(p)
        db.flush()
    lb.peserta_baru = {**pb, "disetujui_waktu": datetime.now().isoformat(timespec="seconds")}
    lb.peserta_baru_status = "disetujui"
    lb.catatan = validate(db, lb)
    db.commit()
    return lembar_out(lb, detail=True)


@router.delete("/lembar/{lid}")
def delete_lembar(lid: int, db: Session = Depends(get_db)):
    lb = _get(db, M.Lembar, lid, "Lembar")
    reading.remove_files(lb)
    db.delete(lb)
    db.commit()
    return {"ok": True}
